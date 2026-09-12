import streamlit as st
import streamlit.components.v1 as components
import fitz  # PyMuPDF
import docx
import json
import io
import requests
from datetime import datetime
from PIL import Image
import pytesseract
from groq import Groq
import firebase_admin
from firebase_admin import credentials, firestore

st.set_page_config(page_title="AI Job Applicant Agent", page_icon="🧑‍💼", layout="centered")

# ---------------- Custom Styling ----------------
st.markdown("""
<style>
    .stApp { background-color: #0e1117; }
    h1, h2, h3 { font-weight: 700; }
    .score-card {
        border-radius: 14px;
        padding: 22px;
        text-align: center;
        margin-bottom: 20px;
    }
    .score-number { font-size: 48px; font-weight: 800; margin: 0; }
    .score-label { font-size: 14px; opacity: 0.85; margin-top: -6px; }
    .skill-chip {
        display: inline-block;
        background-color: #3a1f1f;
        color: #ff8080;
        border: 1px solid #6b2b2b;
        border-radius: 20px;
        padding: 5px 14px;
        margin: 4px 6px 4px 0;
        font-size: 13px;
    }
    .suggestion-card {
        background-color: #1a1d24;
        border-left: 4px solid #4c8bf5;
        border-radius: 8px;
        padding: 10px 16px;
        margin-bottom: 10px;
        font-size: 14px;
    }
    .auth-header {
        text-align: center;
        padding: 10px 0 20px 0;
    }
    .auth-emoji { font-size: 60px; }
    .auth-title { font-size: 30px; font-weight: 800; margin: 4px 0 0 0; }
    .auth-subtitle { font-size: 15px; opacity: 0.75; margin-top: 4px; }
</style>
""", unsafe_allow_html=True)

# ---------------- Firebase Setup ----------------
FIREBASE_WEB_API_KEY = st.secrets["FIREBASE_WEB_API_KEY"]

if not firebase_admin._apps:
    cred_dict = dict(st.secrets["firebase_service_account"])
    cred = credentials.Certificate(cred_dict)
    firebase_admin.initialize_app(cred)

db = firestore.client()

# ---------------- Groq Setup ----------------
try:
    groq_api_key = st.secrets["GROQ_API_KEY"]
except Exception:
    groq_api_key = None

groq_client = Groq(api_key=groq_api_key) if groq_api_key else None


# ---------------- Firebase Auth Helpers ----------------
def firebase_sign_up(email, password):
    url = f"https://identitytoolkit.googleapis.com/v1/accounts:signUp?key={FIREBASE_WEB_API_KEY}"
    payload = {"email": email, "password": password, "returnSecureToken": True}
    r = requests.post(url, data=payload)
    return r.json()


def firebase_sign_in(email, password):
    url = f"https://identitytoolkit.googleapis.com/v1/accounts:signInWithPassword?key={FIREBASE_WEB_API_KEY}"
    payload = {"email": email, "password": password, "returnSecureToken": True}
    r = requests.post(url, data=payload)
    return r.json()


# ---------------- Resume Extraction Helpers ----------------
def ocr_image(image: Image.Image) -> str:
    try:
        return pytesseract.image_to_string(image)
    except Exception:
        return ""


def extract_text_from_pdf(file_bytes: bytes) -> str:
    text_parts = []
    pdf = fitz.open(stream=file_bytes, filetype="pdf")
    for page in pdf:
        page_text = page.get_text().strip()
        if page_text:
            text_parts.append(page_text)
        else:
            pix = page.get_pixmap(dpi=200)
            img_bytes = pix.tobytes("png")
            image = Image.open(io.BytesIO(img_bytes))
            ocr_text = ocr_image(image)
            if ocr_text.strip():
                text_parts.append(ocr_text)
    pdf.close()
    return "\n".join(text_parts)


def extract_text_from_docx(file_bytes: bytes) -> str:
    document = docx.Document(io.BytesIO(file_bytes))
    parts = [para.text for para in document.paragraphs if para.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    try:
        for rel in document.part.rels.values():
            if "image" in rel.reltype:
                image_bytes = rel.target_part.blob
                image = Image.open(io.BytesIO(image_bytes))
                ocr_text = ocr_image(image)
                if ocr_text.strip():
                    parts.append(ocr_text)
    except Exception:
        pass
    return "\n".join(parts)


def extract_text_from_image(file_bytes: bytes) -> str:
    image = Image.open(io.BytesIO(file_bytes))
    return ocr_image(image)


def extract_resume_text(uploaded_file) -> str:
    file_bytes = uploaded_file.read()
    name = uploaded_file.name.lower()
    if name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif name.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    elif name.endswith((".png", ".jpg", ".jpeg", ".webp")):
        return extract_text_from_image(file_bytes)
    else:
        return ""


def analyze_resume(resume_text, job_description):
    prompt = f"""
You are an ATS (Applicant Tracking System) and career expert.
Compare the following RESUME against the JOB DESCRIPTION.

Return ONLY a valid JSON object with these exact keys, no extra text before or after:
{{
  "ats_score": <number from 0 to 100>,
  "missing_skills": [<list of important skills/keywords from the job description that are missing in the resume>],
  "suggestions": [<list of 3-5 short, actionable suggestions to improve the resume for this job>]
}}

RESUME:
{resume_text}

JOB DESCRIPTION:
{job_description}
"""
    response = groq_client.chat.completions.create(
        model="openai/gpt-oss-120b",
        messages=[{"role": "user", "content": prompt}],
        temperature=0.3,
    )
    raw_output = response.choices[0].message.content.strip()
    if raw_output.startswith("```"):
        raw_output = raw_output.strip("`")
        if raw_output.startswith("json"):
            raw_output = raw_output[4:]
        raw_output = raw_output.strip()
    return json.loads(raw_output)


# ---------------- Firestore History Helpers ----------------
def save_analysis_to_history(uid, job_description, result):
    doc = {
        "job_description": job_description[:300],
        "ats_score": result["ats_score"],
        "missing_skills": result["missing_skills"],
        "suggestions": result["suggestions"],
        "timestamp": datetime.utcnow().isoformat(),
    }
    db.collection("users").document(uid).collection("history").add(doc)


def get_user_history(uid):
    docs = (
        db.collection("users")
        .document(uid)
        .collection("history")
        .order_by("timestamp", direction=firestore.Query.DESCENDING)
        .limit(10)
        .stream()
    )
    return [d.to_dict() for d in docs]


# ---------------- UI Helper Widgets ----------------
def score_color(score):
    if score >= 70:
        return "#1f6b3a", "#42d17a"   # bg, text
    elif score >= 40:
        return "#5c4a12", "#f5c542"
    else:
        return "#5c1f1f", "#ff6b6b"


def render_score(score):
    bg, fg = score_color(score)
    st.markdown(f"""
    <div class="score-card" style="background-color:{bg}22; border: 1px solid {bg};">
        <p class="score-number" style="color:{fg};">{score}<span style="font-size:22px;">/100</span></p>
        <p class="score-label" style="color:{fg};">ATS Match Score</p>
    </div>
    """, unsafe_allow_html=True)


def render_skill_chips(skills):
    if not skills:
        st.write("No major missing skills found!")
        return
    chips_html = "".join([f'<span class="skill-chip">{s}</span>' for s in skills])
    st.markdown(chips_html, unsafe_allow_html=True)


def render_suggestions(suggestions):
    for s in suggestions:
        st.markdown(f'<div class="suggestion-card">💡 {s}</div>', unsafe_allow_html=True)


def result_to_text(job_description, result):
    lines = [
        f"ATS Score: {result['ats_score']}/100",
        "",
        "Missing Skills:",
    ]
    lines += [f"- {s}" for s in result["missing_skills"]] or ["None"]
    lines += ["", "Suggestions:"]
    lines += [f"- {s}" for s in result["suggestions"]]
    return "\n".join(lines)


def copy_button(text, key):
    safe_text = json.dumps(text)
    html_code = f"""
    <button id="copybtn_{key}" onclick='navigator.clipboard.writeText({safe_text});
        document.getElementById("copybtn_{key}").innerText="Copied!";
        setTimeout(()=>{{document.getElementById("copybtn_{key}").innerText="📋 Copy";}}, 1500);'
        style="background:#1a1d24;color:#e0e0e0;border:1px solid #3a3d46;padding:8px 16px;
        border-radius:8px;cursor:pointer;font-size:14px;width:100%;">
        📋 Copy
    </button>
    """
    components.html(html_code, height=45)


def read_aloud_button(text, key):
    safe_text = json.dumps(text)
    html_code = f"""
    <button onclick='window.speechSynthesis.cancel(); var u = new SpeechSynthesisUtterance({safe_text}); window.speechSynthesis.speak(u);'
        style="background:#1a1d24;color:#e0e0e0;border:1px solid #3a3d46;padding:8px 16px;
        border-radius:8px;cursor:pointer;font-size:14px;width:100%;">
        🔊 Read Aloud
    </button>
    """
    components.html(html_code, height=45)


# ---------------- Session State ----------------
if "user" not in st.session_state:
    st.session_state.user = None
if "last_resume_text" not in st.session_state:
    st.session_state.last_resume_text = None
if "last_result" not in st.session_state:
    st.session_state.last_result = None
if "last_job_description" not in st.session_state:
    st.session_state.last_job_description = ""


def logout():
    st.session_state.user = None
    st.session_state.last_resume_text = None
    st.session_state.last_result = None
    st.rerun()


def run_analysis(resume_text, job_description, uid):
    with st.spinner("Analyzing your resume..."):
        result = analyze_resume(resume_text, job_description)
    st.session_state.last_resume_text = resume_text
    st.session_state.last_result = result
    st.session_state.last_job_description = job_description
    save_analysis_to_history(uid, job_description, result)


# ---------------- Auth Screens ----------------
def show_auth_screen():
    st.markdown("""
    <div class="auth-header">
        <div class="auth-emoji">🧑‍💼✨</div>
        <p class="auth-title">AI Job Applicant Agent</p>
        <p class="auth-subtitle">Land your next job faster — check your resume against any job in seconds.</p>
    </div>
    """, unsafe_allow_html=True)

    tab_login, tab_signup = st.tabs(["🔐 Login", "🆕 Sign Up"])

    with tab_login:
        email = st.text_input("Email", key="login_email")
        password = st.text_input("Password", type="password", key="login_password")
        if st.button("Login", use_container_width=True):
            if not email or not password:
                st.error("Please enter both email and password.")
            else:
                result = firebase_sign_in(email, password)
                if "idToken" in result:
                    st.session_state.user = {"uid": result["localId"], "email": result["email"]}
                    st.rerun()
                else:
                    error_msg = result.get("error", {}).get("message", "Login failed.")
                    st.error(f"Login failed: {error_msg}")

    with tab_signup:
        new_email = st.text_input("Email", key="signup_email")
        new_password = st.text_input("Password (min 6 characters)", type="password", key="signup_password")
        if st.button("Create Account", use_container_width=True):
            if not new_email or not new_password:
                st.error("Please enter both email and password.")
            elif len(new_password) < 6:
                st.error("Password must be at least 6 characters.")
            else:
                result = firebase_sign_up(new_email, new_password)
                if "idToken" in result:
                    st.session_state.user = {"uid": result["localId"], "email": result["email"]}
                    st.success("Account created successfully!")
                    st.rerun()
                else:
                    error_msg = result.get("error", {}).get("message", "Sign up failed.")
                    st.error(f"Sign up failed: {error_msg}")


# ---------------- Main App ----------------
def show_main_app():
    user = st.session_state.user

    with st.sidebar:
        st.subheader("👤 Account")
        st.write(f"**Email:** {user['email']}")
        st.caption("AI Job Applicant Agent helps you match your resume to any job description, spot missing skills, and get instant improvement tips.")
        if st.button("Logout", use_container_width=True):
            logout()

    st.title("🧑‍💼 AI Job Applicant Agent")
    st.write("Upload your resume and paste a job description. The AI will check your ATS score and missing skills.")

    tab_analyze, tab_history = st.tabs(["Analyze Resume", "History"])

    with tab_analyze:
        uploaded_file = st.file_uploader(
            "Upload your Resume (PDF, DOCX, or Image)",
            type=["pdf", "docx", "png", "jpg", "jpeg", "webp"]
        )
        job_description = st.text_area("Paste the Job Description here", height=220, key="jd_input")

        if st.button("Analyze Resume", use_container_width=True):
            if not groq_client:
                st.error("Groq API key is not configured.")
            elif not uploaded_file:
                st.error("Please upload your resume.")
            elif not job_description.strip():
                st.error("Please paste a job description.")
            else:
                with st.spinner("Reading your resume..."):
                    try:
                        resume_text = extract_resume_text(uploaded_file)
                        if not resume_text.strip():
                            st.error("Could not extract any text from this file. Try a clearer scan or a different file.")
                        else:
                            run_analysis(resume_text, job_description, user["uid"])
                    except json.JSONDecodeError:
                        st.error("The AI response could not be read properly. Please try again.")
                    except Exception as e:
                        st.error(f"Something went wrong: {e}")

        # ---- Show results if available ----
        if st.session_state.last_result:
            result = st.session_state.last_result
            st.divider()
            render_score(result["ats_score"])

            st.subheader("❌ Missing Skills")
            render_skill_chips(result["missing_skills"])

            st.subheader("💡 Suggestions")
            render_suggestions(result["suggestions"])

            st.write("")
            col1, col2, col3, col4 = st.columns(4)
            with col1:
                if st.button("👍 Like", use_container_width=True, key="like_btn"):
                    st.toast("Thanks for your feedback!")
            with col2:
                if st.button("👎 Dislike", use_container_width=True, key="dislike_btn"):
                    st.toast("Thanks — we'll keep improving!")
            with col3:
                copy_button(result_to_text(st.session_state.last_job_description, result), key="result")
            with col4:
                read_aloud_button(result_to_text(st.session_state.last_job_description, result), key="result")

            st.divider()
            with st.expander("✏️ Edit Job Description & Re-analyze (no need to re-upload resume)"):
                edited_jd = st.text_area(
                    "Job Description",
                    value=st.session_state.last_job_description,
                    height=180,
                    key="edit_jd"
                )
                if st.button("Re-analyze with edited description", use_container_width=True):
                    try:
                        run_analysis(st.session_state.last_resume_text, edited_jd, user["uid"])
                        st.rerun()
                    except Exception as e:
                        st.error(f"Something went wrong: {e}")

    with tab_history:
        st.subheader("📜 Your Past Analyses")
        history = get_user_history(user["uid"])
        if not history:
            st.write("No history yet. Analyze a resume to see it here!")
        else:
            for item in history:
                with st.expander(f"Score: {item['ats_score']}/100 — {item['timestamp'][:19]}"):
                    st.write(f"**Job Description (excerpt):** {item['job_description']}")
                    st.write("**Missing Skills:**")
                    render_skill_chips(item["missing_skills"])
                    st.write("**Suggestions:**")
                    render_suggestions(item["suggestions"])


# ---------------- Router ----------------
if st.session_state.user is None:
    show_auth_screen()
else:
    show_main_app()